from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import List

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from .. import schemas
from ..database import get_db
from ..deps import get_current_user, require_role
from ..models import DataRecord, DataUpload, User

router = APIRouter(prefix="/uploads", tags=["uploads"])


REQUIRED_HEADERS = {"Ad", "Soyad", "Telefon", "Kategori"}


def _read_excel(file_bytes: bytes) -> List[dict]:
    workbook = load_workbook(BytesIO(file_bytes))
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    if not headers:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Excel file must have headers")

    header_map = {header: idx for idx, header in enumerate(headers)}
    if not REQUIRED_HEADERS.issubset(header_map.keys()):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Excel headers must include: {', '.join(REQUIRED_HEADERS)}",
        )

    records = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        record = {
            "first_name": str(row[header_map["Ad"]]).strip(),
            "last_name": str(row[header_map["Soyad"]]).strip(),
            "phone": str(row[header_map["Telefon"]]).strip(),
            "category": str(row[header_map["Kategori"]]).strip() if row[header_map["Kategori"]] else None,
        }
        records.append(record)
    return records


@router.post("/excel", response_model=schemas.DataUploadResponse, status_code=status.HTTP_201_CREATED)
async def upload_excel(
    title: str = Form(...),
    price_per_record: Decimal = Form(...),
    description: str | None = Form(None),
    file: UploadFile = File(...),
    current_user: User = Depends(require_role(["user"])),
    db: Session = Depends(get_db),
):
    file_bytes = await file.read()
    records = _read_excel(file_bytes)
    if not records:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Excel contains no data")

    upload = DataUpload(
        user_id=current_user.id,
        title=title,
        description=description,
        price_per_record=price_per_record,
        total_records=len(records),
        created_at=datetime.utcnow(),
    )

    for record in records:
        upload.records.append(DataRecord(**record))

    db.add(upload)
    db.commit()
    db.refresh(upload)
    return upload


@router.get("/", response_model=list[schemas.DataUploadResponse])
def list_uploads(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(DataUpload)
    if current_user.role == "user":
        query = query.filter(DataUpload.user_id == current_user.id)
    elif current_user.role == "operator":
        query = query.filter(DataUpload.is_approved.is_(True))
    return query.all()
